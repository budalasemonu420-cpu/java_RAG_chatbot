import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from core.documents import (
    DocumentProcessingError,
    ExtractedPage,
    chunk_document,
    clean_text,
    extract_document,
)


class DocumentTextTests(unittest.TestCase):
    def test_clean_text_removes_controls_and_preserves_paragraphs(self):
        self.assertEqual(clean_text("  First\nwrapped\x00 line.\n\n Second   section "),
                         "First wrapped line.\n\nSecond section")

    def test_chunk_document_keeps_source_and_page_metadata(self):
        pages = [ExtractedPage(3, "word " * 110)]
        chunks = chunk_document(pages, "syllabus.pdf", chunk_size=100, overlap=20)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.source == "syllabus.pdf" for chunk in chunks))
        self.assertTrue(all(chunk.page_number == 3 for chunk in chunks))
        self.assertEqual(chunks[0].chunk_id, "syllabus.pdf:3:1")
        self.assertTrue(all(len(chunk.text) <= 100 for chunk in chunks))

    def test_docx_extraction_includes_paragraphs_and_tables(self):
        from docx import Document

        with TemporaryDirectory() as directory:
            path = Path(directory) / "course.docx"
            document = Document()
            document.add_paragraph("Attendance is required.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Assessment"
            table.cell(0, 1).text = "Final exam"
            document.save(path)
            pages = extract_document(path)
        self.assertIsNone(pages[0].page_number)
        self.assertIn("Attendance is required", pages[0].text)
        self.assertIn("Final exam", pages[0].text)

    def test_pdf_extraction_preserves_page_number(self):
        import pymupdf

        with TemporaryDirectory() as directory:
            path = Path(directory) / "notes.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "Eligibility requires a completed application.")
            document.save(path)
            document.close()
            pages = extract_document(path)
        self.assertEqual(pages[0].page_number, 1)
        self.assertIn("completed application", pages[0].text)

    def test_empty_or_unsupported_document_fails_clearly(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "empty.docx"
            from docx import Document

            Document().save(path)
            with self.assertRaisesRegex(DocumentProcessingError, "No selectable text"):
                extract_document(path)
            with self.assertRaisesRegex(DocumentProcessingError, "Upload a PDF or DOCX"):
                extract_document(Path(directory) / "notes.txt")


if __name__ == "__main__":
    unittest.main()