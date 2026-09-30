"""Local PDF/DOCX extraction, conservative text cleaning, and chunking."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re


class DocumentProcessingError(Exception):
    """Raised when a supported document cannot be read or has no text."""


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int | None
    text: str


@dataclass(frozen=True)
class DocumentChunk:
    chunk_id: str
    source: str
    page_number: int | None
    text: str


def clean_text(text: str) -> str:
    """Normalize whitespace while preserving paragraph boundaries and meaning."""
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    paragraphs: list[str] = []
    current_lines: list[str] = []

    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        normalized = re.sub(r"\s+", " ", line).strip()
        if normalized:
            current_lines.append(normalized)
        elif current_lines:
            paragraphs.append(" ".join(current_lines))
            current_lines = []

    if current_lines:
        paragraphs.append(" ".join(current_lines))
    return "\n\n".join(paragraphs)


def extract_document(path: str | Path) -> list[ExtractedPage]:
    """Extract text from a PDF or DOCX without making optional imports mandatory."""
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            import pymupdf

            with pymupdf.open(path) as pdf:
                pages = [
                    ExtractedPage(index, clean_text(page.get_text("text")))
                    for index, page in enumerate(pdf, start=1)
                ]
        elif suffix == ".docx":
            from docx import Document

            document = Document(path)
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            pages = [ExtractedPage(None, clean_text("\n\n".join(parts)))]
        else:
            raise DocumentProcessingError("Upload a PDF or DOCX file.")
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError(f"Could not read {path.name}: {exc}") from exc

    pages = [page for page in pages if page.text]
    if not pages:
        raise DocumentProcessingError(
            "No selectable text was found. This may be an image-only PDF; run OCR first."
        )
    return pages


def _text_windows(text: str, chunk_size: int, overlap: int) -> list[str]:
    windows: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + chunk_size // 2, end)
            if boundary > start:
                end = boundary
        value = text[start:end].strip()
        if value:
            windows.append(value)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
        while start < len(text) and text[start].isspace():
            start += 1
    return windows


def chunk_document(
    pages: list[ExtractedPage],
    source: str,
    chunk_size: int = 900,
    overlap: int = 120,
) -> list[DocumentChunk]:
    """Split extracted pages into character-bounded chunks with source metadata."""
    if chunk_size < 100:
        raise ValueError("chunk_size must be at least 100 characters")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be between 0 and chunk_size - 1")

    chunks: list[DocumentChunk] = []
    for page in pages:
        text = clean_text(page.text)
        for window in _text_windows(text, chunk_size, overlap):
            chunk_id = f"{Path(source).name}:{page.page_number or 'docx'}:{len(chunks) + 1}"
            chunks.append(DocumentChunk(chunk_id, source, page.page_number, window))
    return chunks